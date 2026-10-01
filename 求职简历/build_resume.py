from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path(__file__).parent
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5)
sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
 s.font.size=Pt(10)
 s.paragraph_format.space_after=Pt(4)
 s.paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(24)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 2'].font.size=Pt(10.5)
doc.styles['Heading 2'].font.bold=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(5)
def p(t,style=None):
 return doc.add_paragraph(t,style)
def bullet(label,text):
 x=p(''); x.paragraph_format.left_indent=Cm(.25); x.paragraph_format.first_line_indent=Cm(-.25)
 x.add_run('• '+label+'：').bold=True; x.add_run(text)
p('姓名待补充','Title')
p('求职意向：Java 后端开发 / 软件开发    |    2027 届本科毕业生')
p('21 岁    |    籍贯：内蒙古赤峰    |    18648100212    |    1206425793@qq.com')
p('教育背景','Heading 1')
x=p('东北大学    |    物联网工程    |    本科'); x.runs[0].bold=True
p('预计 2027 年毕业')
p('专业技能','Heading 1')
bullet('后端开发','使用 Java、Spring Boot、Spring MVC 开发业务接口，结合 Spring Validation 进行参数校验。')
bullet('数据库与缓存','使用 MySQL、Redis 进行数据存储与缓存；项目涉及 Redis List、Set、原子操作及 Redisson 分布式锁。')
bullet('消息与并发','使用 Kafka 处理异步业务，结合 Caffeine 缓存、消息幂等和限流机制优化业务处理流程。')
p('项目经历','Heading 1')
p('直播互动平台后端    |    Java 项目实践    |    时间待补充','Heading 2')
p('技术栈：Spring Boot / Spring MVC / MySQL / Redis / Kafka / Caffeine / Redisson')
p('项目简介：围绕直播间管理、礼物互动、红包及消息推送等场景，开发后端业务接口。')
bullet('接口与房间管理','实现登录态校验和接口访问控制；开发开播、关播接口，处理房间状态更新、直播记录归档及缓存清理。')
bullet('礼物与红包','通过 Kafka 异步处理送礼消息；基于 Redis List 管理红包池，通过原子弹出分配红包，异步处理领取记录与入账流程。')
bullet('消息推送','使用 Redis Set 维护房间在线用户，按连接所在机器分组批量推送消息，减少重复远程调用。')
bullet('缓存与幂等','使用 Caffeine 缓存礼物及红包配置；结合消息唯一标识与 Redis 去重，处理重复消费问题。')
p('通用红包雨 SDK    |    Java 项目实践    |    时间待补充','Heading 2')
p('技术栈：Spring Boot / MySQL / Redis / Redisson / Spring Validation / 消息队列')
p('项目简介：封装红包活动管理、红包分配与领取能力，为业务接入提供统一接口。')
bullet('活动与分配','实现活动创建、启停和配置查询；支持固定金额及随机金额红包，采用二倍均值法拆分随机红包。')
bullet('并发与防重','使用 Redis 预存库存并执行原子扣减；结合分布式锁与用户、活动联合唯一索引，控制重复领取和并发冲突。')
bullet('风控与封装','围绕用户、设备、IP 维度实现请求频率控制与资格校验；封装创建活动、抢红包、查询记录等接口，降低接入耦合。')
p('校园经历','Heading 1')
p('计算机科学与工程学院文化艺术中心    |    文艺部部长','Heading 2')
p('2024.09—2025.09')
bullet('活动筹办','参与筹办院级元旦晚会等文艺活动，承担文艺部相关组织工作。')
doc.core_properties.title='2027届Java后端开发校招简历'
doc.core_properties.author=''
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):
  border.getparent().remove(border)
for para in doc.paragraphs:
 pf=para.paragraph_format
 pf.line_spacing=Pt(14)
 pf.space_after=Pt(3)
 pp=para._p.get_or_add_pPr()
 snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); pp.append(snap)
 for border in list(pp.findall(qn('w:pBdr'))): pp.remove(border)
doc.save(out/'Java后端开发_校招简历_初稿.docx')
