from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out=Path(__file__).parent
doc=Document(out/'Java后端开发_校招简历_初稿.docx')
section=doc.sections[0]
section.left_margin=section.right_margin=Cm(1.4)
section.top_margin=section.bottom_margin=Cm(1.3)
skills=[]; active=False
for p in doc.paragraphs:
 if p.text=='专业技能': active=True
 if p.text=='项目经历': active=False
 if active: skills.append(p._p)
body=doc._element.body
for el in skills:
 body.remove(el); body.insert(len(body)-1,el)
for p in doc.paragraphs:
 text=p.text
 p.paragraph_format.line_spacing=Pt(15)
 p.paragraph_format.space_after=Pt(4)
 if p.style.name=='Heading 1':
  p.paragraph_format.space_before=Pt(9)
  p.paragraph_format.space_after=Pt(7)
  for r in p.runs: r.font.color.rgb=RGBColor.from_string('345361')
  borders=OxmlElement('w:pBdr'); bottom=OxmlElement('w:bottom')
  for key,value in [('val','single'),('sz','6'),('space','4'),('color','56707C')]: bottom.set(qn('w:'+key),value)
  borders.append(bottom); p._p.get_or_add_pPr().append(borders)
 if text=='东北大学    |    物联网工程    |    本科':
  p.text='东北大学  |  物联网工程  |  本科\t2027 年毕业'
  p.runs[0].bold=True
 if text=='预计 2027 年毕业': p._element.getparent().remove(p._element); continue
 if 'Java 项目实践' in text:
  p.text=text.split('    |')[0]+'  |  项目实践\t时间待补充'
  p.runs[0].bold=True
 if text=='计算机科学与工程学院文化艺术中心    |    文艺部部长':
  p.text='计算机科学与工程学院文化艺术中心  |  文艺部部长\t2024.09—2025.09'
  p.runs[0].bold=True
 if text=='2024.09—2025.09': p._element.getparent().remove(p._element); continue
 p.paragraph_format.tab_stops.add_tab_stop(Cm(18.2),WD_TAB_ALIGNMENT.RIGHT)
 if text.startswith('技术栈：') or text.startswith('项目简介：'):
  label,content=text.split('：',1); p.clear(); p.add_run(label+'：').bold=True; p.add_run(content)
 if p.style.name not in ['Title','Heading 1','Heading 2']:
  for r in p.runs: r.font.size=Pt(10.5)
 if p.style.name=='Title': p.paragraph_format.line_spacing=Pt(29)
doc.save(out/'Java后端开发_校招简历_参考版.docx')
