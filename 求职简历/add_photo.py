from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).parent
doc=Document(root/'乌凡杰_Java后端开发_校招简历.docx')
photo=r'E:/Temp/WeChat/WeChat Files/wxid_yt5np8zckcfn22/FileStorage/Temp/a03669fd9bee6e4500e2039ee667ac1.png'
shape=doc.paragraphs[0].add_run().add_picture(photo,width=Cm(2.4))
inline=shape._inline
anchor=OxmlElement('wp:anchor')
for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'251658240','behindDoc':'0','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items(): anchor.set(k,v)
pos=OxmlElement('wp:simplePos'); pos.set('x','0'); pos.set('y','0'); anchor.append(pos)
for axis,offset in [('H',15.8),('V',0)]:
 p=OxmlElement('wp:position'+axis); p.set('relativeFrom','margin')
 off=OxmlElement('wp:posOffset'); off.text=str(int(Cm(offset))); p.append(off); anchor.append(p)
anchor.append(inline.find(qn('wp:extent')))
effect=OxmlElement('wp:effectExtent')
for k in ['l','t','r','b']: effect.set(k,'0')
anchor.append(effect); anchor.append(OxmlElement('wp:wrapNone'))
for tag in ['wp:docPr','wp:cNvGraphicFramePr','a:graphic']:
 el=inline.find(qn(tag))
 if el is not None: anchor.append(el)
anchor.find(qn('wp:docPr')).set('descr','乌凡杰证件照')
inline.getparent().replace(inline,anchor)
# Reserve space below the contact block for the full-height portrait.
doc.paragraphs[2].paragraph_format.space_after=Pt(30)
doc.save(root/'乌凡杰_Java后端开发_校招简历_含照片.docx')
