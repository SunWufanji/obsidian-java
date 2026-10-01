from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

root=Path(__file__).parent
d=Document(root/'乌凡杰_Java后端开发_校招简历_含照片.docx')
for p in d.paragraphs:
 if p.text.startswith('东北大学'):
  p.text='东北大学 985 211 双一流  ·  物联网工程  ·  本科\t2027 年毕业'
  p.runs[0].bold=True
  p.runs[0].font.size=Pt(10.5)
  el=OxmlElement('w:p'); p._p.addnext(el)
  np=Paragraph(el,p._parent)
  np.paragraph_format.line_spacing=Pt(15)
  np.paragraph_format.space_after=Pt(4)
  np.add_run('主修课程：').bold=True
  np.add_run('待补充实际修读课程')
  for r in np.runs: r.font.size=Pt(10.5)
  break
d.save(root/'乌凡杰_Java后端开发_校招简历_教育更新版.docx')
